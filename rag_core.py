"""
RAG pipeline with LLM-based retrieval self-assessment.

Flow: retrieve -> grade chunks -> if sufficient context, answer;
      else rewrite query -> re-retrieve (up to MAX_RETRIEVAL_ITERATIONS).

Differences from the terminal version (rag_pipeline.py):
  * LLM provider is pluggable: "groq" (hosted, free tier) or "ollama" (local).
  * Grading is batched into ONE call per iteration instead of one call per chunk,
    cutting requests-per-query from ~15 to ~3 so free-tier rate limits hold up.
  * Embeddings use Chroma's bundled ONNX MiniLM-L6-v2 (same model as before)
    instead of sentence-transformers, which removes the PyTorch dependency and
    keeps memory inside free-host limits.
  * No printing. The pipeline yields structured trace events so any UI can
    render the reasoning loop.
"""

from __future__ import annotations

import hashlib
import json
import logging
import os
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterator, Literal

import chromadb
from chromadb.utils import embedding_functions
from pypdf import PdfReader

# pypdf logs harmless "Ignoring wrong pointing object" warnings for some PDFs.
logging.getLogger("pypdf").setLevel(logging.ERROR)

# ── LLM configuration ────────────────────────────────────────────────────────
# "groq"   -> hosted inference, needs GROQ_API_KEY (free tier, no card)
# "ollama" -> local model, for running on your own machine
LLM_PROVIDER = os.getenv("LLM_PROVIDER", "groq").lower()

# Groq model IDs and per-account access vary. Check https://console.groq.com/docs/models
# if you get a 404/model_not_found and update GROQ_MODEL in your env/secrets.
GROQ_MODEL = os.getenv("GROQ_MODEL", "openai/gpt-oss-120b")

# Reasoning models (gpt-oss) think before answering, and that thinking counts
# against the token limit. This headroom stops short calls (grading, rewriting)
# from running out of tokens before they produce an answer.
REASONING_HEADROOM = 1024
OLLAMA_MODEL = os.getenv("OLLAMA_MODEL", "llama3.2")

# ── Pipeline ─────────────────────────────────────────────────────────────────
MAX_RETRIEVAL_ITERATIONS = int(os.getenv("MAX_RETRIEVAL_ITERATIONS", "3"))
MIN_RELEVANT_CHUNKS = int(os.getenv("MIN_RELEVANT_CHUNKS", "1"))
TOP_K = int(os.getenv("TOP_K", "5"))

# ── ChromaDB ─────────────────────────────────────────────────────────────────
CHROMA_PATH = os.getenv("CHROMA_PATH", "./chroma_db")
COLLECTION_NAME = "rag_docs_v2"

# ── Document ingestion ───────────────────────────────────────────────────────
DOCS_DIR = os.getenv("DOCS_DIR", "./demo_docs")
CHUNK_SIZE = 1000
CHUNK_OVERLAP = 200
DOC_INDEX_PATH = os.getenv("DOC_INDEX_PATH", "./doc_index.json")

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".csv", ".xlsx", ".xls"}


# ---------------------------------------------------------------------------
# LLM adapter
# ---------------------------------------------------------------------------

class LLMError(RuntimeError):
    """Raised when the configured provider cannot be reached or is misconfigured."""


def _groq_client():
    from groq import Groq

    api_key = os.getenv("GROQ_API_KEY")
    if not api_key:
        raise LLMError(
            "GROQ_API_KEY is not set. Add it to .streamlit/secrets.toml locally, "
            "or to Settings -> Secrets in Streamlit Community Cloud."
        )
    return Groq(api_key=api_key)


def _is_reasoning_model(model: str) -> bool:
    return model.startswith("openai/gpt-oss")


def _groq_params(max_tokens: int | None) -> dict:
    """Model-specific request settings for Groq."""
    params: dict = {"model": GROQ_MODEL}
    if _is_reasoning_model(GROQ_MODEL):
        # Keep thinking short and out of the reply; leave room for it in the budget.
        params["reasoning_effort"] = "low"
        params["include_reasoning"] = False
        if max_tokens:
            params["max_completion_tokens"] = max_tokens + REASONING_HEADROOM
    elif max_tokens:
        params["max_completion_tokens"] = max_tokens
    return params


def chat(messages: list[dict], temperature: float = 0.0, max_tokens: int = 1024) -> str:
    """Single-shot completion. Returns the assistant's text."""
    if LLM_PROVIDER == "groq":
        client = _groq_client()
        response = client.chat.completions.create(
            messages=messages,
            temperature=temperature,
            **_groq_params(max_tokens),
        )
        return (response.choices[0].message.content or "").strip()

    if LLM_PROVIDER == "ollama":
        import ollama

        response = ollama.chat(model=OLLAMA_MODEL, messages=messages)
        return response.message.content.strip()

    raise LLMError(f"Unknown LLM_PROVIDER: {LLM_PROVIDER!r}. Use 'groq' or 'ollama'.")


def chat_stream(messages: list[dict], temperature: float = 0.2) -> Iterator[str]:
    """Streaming completion. Yields text tokens."""
    if LLM_PROVIDER == "groq":
        client = _groq_client()
        stream = client.chat.completions.create(
            messages=messages,
            temperature=temperature,
            stream=True,
            **_groq_params(None),
        )
        for chunk in stream:
            token = chunk.choices[0].delta.content
            if token:
                yield token
        return

    if LLM_PROVIDER == "ollama":
        import ollama

        for chunk in ollama.chat(model=OLLAMA_MODEL, messages=messages, stream=True):
            token = chunk.message.content
            if token:
                yield token
        return

    raise LLMError(f"Unknown LLM_PROVIDER: {LLM_PROVIDER!r}. Use 'groq' or 'ollama'.")


def active_model_name() -> str:
    return GROQ_MODEL if LLM_PROVIDER == "groq" else OLLAMA_MODEL


# ---------------------------------------------------------------------------
# Text cleaning and chunking
# ---------------------------------------------------------------------------

def _clean(text: str) -> str:
    """Normalize whitespace while preserving sentence structure."""
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    text = re.sub(r"(?<!\n)\n(?!\n)", " ", text)
    return text.strip()


def chunk_text(text: str, chunk_size: int = CHUNK_SIZE, overlap: int = CHUNK_OVERLAP) -> list[str]:
    """Split text into overlapping chunks that always break at sentence boundaries."""
    sentences = re.split(r"(?<=[.!?])\s+", text)

    chunks: list[str] = []
    current: list[str] = []
    current_len = 0

    for sentence in sentences:
        s_len = len(sentence)
        if current_len + s_len > chunk_size and current:
            chunks.append(" ".join(current))
            tail, tail_len = [], 0
            for s in reversed(current):
                if tail_len + len(s) <= overlap:
                    tail.insert(0, s)
                    tail_len += len(s)
                else:
                    break
            current, current_len = tail, tail_len
        current.append(sentence)
        current_len += s_len

    if current:
        chunks.append(" ".join(current))

    return [c for c in chunks if len(c) > 80]


def _make_records(path: Path, sections: list[tuple[int, str]]) -> list[dict]:
    """Convert (section_num, text) pairs into chunk record dicts."""
    records = []
    for section_num, text in sections:
        cleaned = _clean(text)
        if not cleaned:
            continue
        for chunk_idx, chunk in enumerate(chunk_text(cleaned)):
            uid = hashlib.md5(f"{path.name}|{section_num}|{chunk_idx}".encode()).hexdigest()
            records.append({
                "id": uid,
                "text": chunk,
                "metadata": {"source": path.name, "page": section_num, "chunk": chunk_idx},
            })
    return records


# ---------------------------------------------------------------------------
# Loaders
# ---------------------------------------------------------------------------

def load_pdf(path: Path) -> list[dict]:
    reader = PdfReader(str(path))
    sections = [
        (page_num, page.extract_text() or "")
        for page_num, page in enumerate(reader.pages, start=1)
    ]
    return _make_records(path, sections)


def load_docx(path: Path) -> list[dict]:
    from docx import Document

    doc = Document(str(path))
    lines: list[str] = []
    for para in doc.paragraphs:
        if para.text.strip():
            lines.append(para.text)
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(c.text.strip() for c in row.cells if c.text.strip())
            if row_text:
                lines.append(row_text)
    return _make_records(path, [(1, "\n".join(lines))])


def load_txt(path: Path) -> list[dict]:
    text = path.read_text(encoding="utf-8", errors="replace")
    return _make_records(path, [(1, text)])


def load_excel(path: Path) -> list[dict]:
    import pandas as pd

    xl = pd.ExcelFile(str(path))
    sections: list[tuple[int, str]] = []
    for sheet_idx, sheet_name in enumerate(xl.sheet_names, start=1):
        df = xl.parse(sheet_name).fillna("")
        header = " | ".join(str(c) for c in df.columns)
        rows = [f"Sheet: {sheet_name} | Columns: {header}"]
        for _, row in df.iterrows():
            row_text = " | ".join(
                f"{col}: {val}" for col, val in row.items() if str(val).strip()
            )
            if row_text:
                rows.append(row_text)
        sections.append((sheet_idx, "\n".join(rows)))
    return _make_records(path, sections)


def load_csv(path: Path) -> list[dict]:
    import pandas as pd

    df = pd.read_csv(str(path)).fillna("")
    header = " | ".join(str(c) for c in df.columns)
    rows = [f"Columns: {header}"]
    for _, row in df.iterrows():
        row_text = " | ".join(
            f"{col}: {val}" for col, val in row.items() if str(val).strip()
        )
        if row_text:
            rows.append(row_text)
    return _make_records(path, [(1, "\n".join(rows))])


def load_file(path: Path) -> list[dict]:
    ext = path.suffix.lower()
    if ext == ".pdf":
        return load_pdf(path)
    if ext == ".docx":
        return load_docx(path)
    if ext == ".txt":
        return load_txt(path)
    if ext in {".xlsx", ".xls"}:
        return load_excel(path)
    if ext == ".csv":
        return load_csv(path)
    return []


# ---------------------------------------------------------------------------
# ChromaDB setup
# ---------------------------------------------------------------------------

def _file_hash(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for block in iter(lambda: f.read(65536), b""):
            h.update(block)
    return h.hexdigest()


def _load_doc_index() -> dict[str, str]:
    index_path = Path(DOC_INDEX_PATH)
    if index_path.exists():
        try:
            return json.loads(index_path.read_text())
        except json.JSONDecodeError:
            return {}
    return {}


def _save_doc_index(index: dict[str, str]) -> None:
    try:
        Path(DOC_INDEX_PATH).write_text(json.dumps(index, indent=2))
    except OSError:
        # Read-only filesystem on some hosts: indexing still worked, we just
        # can't cache the hashes. Next cold start re-indexes.
        pass


def _index_records(collection, records: list[dict]) -> None:
    batch_size = 500
    for i in range(0, len(records), batch_size):
        batch = records[i: i + batch_size]
        collection.add(
            ids=[r["id"] for r in batch],
            documents=[r["text"] for r in batch],
            metadatas=[r["metadata"] for r in batch],
        )


@dataclass
class IndexStatus:
    new: int = 0
    updated: int = 0
    unchanged: int = 0
    total_chunks: int = 0
    files: list[str] = field(default_factory=list)


def setup_chroma() -> tuple[object, IndexStatus]:
    """
    Initialise a persistent ChromaDB collection and sync it with DOCS_DIR.

    Uses Chroma's bundled ONNX all-MiniLM-L6-v2 embedder: same model as the
    sentence-transformers version, without pulling PyTorch into the image.
    """
    embedding_fn = embedding_functions.DefaultEmbeddingFunction()
    db = chromadb.PersistentClient(path=CHROMA_PATH)
    collection = db.get_or_create_collection(
        name=COLLECTION_NAME,
        embedding_function=embedding_fn,
        metadata={"hnsw:space": "cosine"},
    )

    docs_path = Path(DOCS_DIR)
    if not docs_path.exists():
        raise FileNotFoundError(f"Documents folder '{DOCS_DIR}' does not exist.")

    doc_files = sorted(
        f for f in docs_path.iterdir()
        if f.is_file() and f.suffix.lower() in SUPPORTED_EXTENSIONS
    )
    if not doc_files:
        exts = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        raise FileNotFoundError(
            f"No supported files found in '{DOCS_DIR}'. Supported: {exts}"
        )

    stored_index = _load_doc_index()
    updated_index = dict(stored_index)
    status = IndexStatus(files=[f.name for f in doc_files])

    for doc_path in doc_files:
        current_hash = _file_hash(doc_path)
        stored_hash = stored_index.get(doc_path.name)

        if stored_hash == current_hash and collection.count() > 0:
            status.unchanged += 1
            continue

        records = load_file(doc_path)
        if not records:
            continue

        if stored_hash is None:
            _index_records(collection, records)
            status.new += 1
        else:
            collection.delete(where={"source": doc_path.name})
            _index_records(collection, records)
            status.updated += 1

        updated_index[doc_path.name] = current_hash

    _save_doc_index(updated_index)
    status.total_chunks = collection.count()
    return collection, status


# ---------------------------------------------------------------------------
# Visitor uploads: a temporary, in-memory index per browser session
# ---------------------------------------------------------------------------

MAX_UPLOAD_FILES = 5
MAX_UPLOAD_MB = 10


def build_upload_collection(files: list[tuple[str, bytes]], session_id: str):
    """
    Index visitor-uploaded files into an in-memory collection that lives only
    for that browser session. Nothing is written to the persistent index.

    Args:
        files:      [(filename, raw_bytes), ...]
        session_id: unique per browser session; used to isolate collections.

    Returns:
        (collection, IndexStatus)
    """
    import tempfile

    if len(files) > MAX_UPLOAD_FILES:
        raise ValueError(f"Please upload at most {MAX_UPLOAD_FILES} files.")

    client = chromadb.EphemeralClient()
    name = f"upload_{session_id}"
    try:
        client.delete_collection(name)  # replace any previous upload set
    except Exception:  # noqa: BLE001 - collection may not exist yet
        pass
    collection = client.create_collection(
        name=name,
        embedding_function=embedding_functions.DefaultEmbeddingFunction(),
        metadata={"hnsw:space": "cosine"},
    )

    status = IndexStatus()
    with tempfile.TemporaryDirectory() as tmp:
        for filename, data in files:
            safe_name = Path(filename).name
            if Path(safe_name).suffix.lower() not in SUPPORTED_EXTENSIONS:
                continue
            if len(data) > MAX_UPLOAD_MB * 1024 * 1024:
                raise ValueError(f"{safe_name} is larger than {MAX_UPLOAD_MB} MB.")
            path = Path(tmp) / safe_name
            path.write_bytes(data)
            records = load_file(path)
            if records:
                _index_records(collection, records)
                status.new += 1
                status.files.append(safe_name)

    if status.new == 0:
        raise ValueError("None of the uploaded files contained readable text.")

    status.total_chunks = collection.count()
    return collection, status


_ACRONYMS = {
    "rag": "RAG", "ai": "AI", "iot": "IoT", "llm": "LLM", "llms": "LLMs",
    "nlp": "NLP", "ml": "ML", "sgd": "SGD", "evor": "EVOR", "fair": "FAIR",
}
_SMALL_WORDS = {"a", "an", "and", "as", "at", "by", "for", "from", "in", "of",
                "on", "or", "the", "to", "with", "vs"}


def display_name(filename: str) -> str:
    """
    Turn a file name into a readable title, e.g.
    '2304.04780v1_a_review_on_explainable_ai.pdf' -> 'A Review on Explainable AI'.
    Names that already look like titles are left alone.
    """
    stem = Path(filename).stem
    stem = re.sub(r"^\d{4}\.\d{4,5}v\d+_", "", stem)          # arXiv id
    stem = re.sub(r"^[a-z]+_\d{7}v\d+_", "", stem)             # old arXiv id
    stem = re.sub(r"^novel_\d+_", "", stem)                     # Gutenberg id
    if "_" not in stem and not stem.islower():
        return stem  # already a readable title, e.g. "Attention Is All You Need"
    words = [w for w in re.split(r"[_\s]+", stem) if w]
    titled = []
    for i, word in enumerate(words):
        if word in _ACRONYMS:
            titled.append(_ACRONYMS[word])
        elif i > 0 and word in _SMALL_WORDS:
            titled.append(word)
        else:
            titled.append(word.capitalize())
    return " ".join(titled)


def retrieve(
    collection, query: str, top_k: int = TOP_K, source: str | None = None
) -> list[dict]:
    """
    Return [{text, source, title, page, distance}] for the top_k nearest chunks.

    If `source` (a file name) is given, only that document is searched.
    """
    kwargs = {"query_texts": [query], "n_results": top_k}
    if source:
        kwargs["where"] = {"source": source}
    results = collection.query(**kwargs)
    documents = results["documents"][0]
    metadatas = results["metadatas"][0]
    distances = results.get("distances", [[None] * len(documents)])[0]
    return [
        {
            "text": doc,
            "source": meta.get("source", "unknown"),
            "title": display_name(meta.get("source", "unknown")),
            "page": meta.get("page", "-"),
            "distance": dist,
        }
        for doc, meta, dist in zip(documents, metadatas, distances)
    ]


# ---------------------------------------------------------------------------
# Document-level tools (used by the agent modes)
# ---------------------------------------------------------------------------

def list_documents(collection) -> list[dict]:
    """
    Every document in the collection as [{file, title, chunks}], sorted by title.
    Lets an agent see what's in the library before deciding where to search.
    """
    metadatas = collection.get(include=["metadatas"])["metadatas"] or []
    counts: dict[str, int] = {}
    for meta in metadatas:
        name = (meta or {}).get("source", "unknown")
        counts[name] = counts.get(name, 0) + 1
    docs = [
        {"file": name, "title": display_name(name), "chunks": n}
        for name, n in counts.items()
    ]
    return sorted(docs, key=lambda d: d["title"].lower())


def _words(text: str) -> set[str]:
    return set(re.findall(r"[a-z0-9]+", text.lower())) - _SMALL_WORDS


def find_document(collection, name: str) -> str | None:
    """
    Match a document name written by an agent (a title, a file name, or a close
    paraphrase like "attention is all you need paper") to an indexed file name.
    Returns None when nothing matches well enough, so the agent can be told so.
    """
    import difflib

    docs = list_documents(collection)
    query = name.strip().lower()
    if not query or not docs:
        return None

    # 1. Exact title or file name.
    for d in docs:
        if query in (d["title"].lower(), d["file"].lower(), Path(d["file"]).stem.lower()):
            return d["file"]

    # 2. Close spelling of a title.
    by_title = {d["title"].lower(): d["file"] for d in docs}
    close = difflib.get_close_matches(query, list(by_title), n=1, cutoff=0.6)
    if close:
        return by_title[close[0]]

    # 3. Best word overlap, if it covers most of what the agent wrote.
    query_words = _words(query) - {"paper", "document", "book", "pdf", "file", "novel"}
    if not query_words:
        return None
    best, best_score = None, 0.0
    for d in docs:
        overlap = len(query_words & _words(d["title"]))
        score = overlap / len(query_words)
        if score > best_score:
            best, best_score = d, score
    if best and best_score >= 0.6:
        return best["file"]
    return None


# ---------------------------------------------------------------------------
# Self-assessment: grading and rewriting
# ---------------------------------------------------------------------------

GRADER_SYSTEM = (
    "You are a strict relevance grader for a retrieval system. You are given a user "
    "query and several numbered document chunks. For each chunk, decide whether it "
    "contains information that genuinely helps answer the query. Be strict: topical "
    "overlap alone is not relevance.\n\n"
    'Reply with JSON only, in exactly this shape: {"relevant": [0, 2]} — a list of the '
    "indices of the relevant chunks. Return an empty list if none are relevant. "
    "No prose, no code fences."
)


def _parse_relevant_indices(raw: str, chunk_count: int) -> tuple[list[int] | None, str | None]:
    """Parse the grader's JSON. Returns (indices, warning)."""
    cleaned = re.sub(r"^```(?:json)?|```$", "", raw.strip(), flags=re.MULTILINE).strip()
    match = re.search(r"\{.*\}", cleaned, flags=re.DOTALL)
    if match:
        cleaned = match.group(0)

    try:
        payload = json.loads(cleaned)
        indices = payload["relevant"]
        valid = sorted({int(i) for i in indices if 0 <= int(i) < chunk_count})
        return valid, None
    except (json.JSONDecodeError, KeyError, TypeError, ValueError):
        return None, "Grader returned unparseable output; treating all chunks as relevant."


def grade_chunks(query: str, chunks: list[dict]) -> tuple[list[int], str | None]:
    """
    Grade all retrieved chunks in a SINGLE LLM call.

    Grading one chunk per call would cost 15+ requests for a 5-chunk / 3-iteration
    query. Batching keeps a full query at ~3 requests,
    which is what makes free-tier rate limits workable for a public demo.
    """
    numbered = "\n\n".join(
        f"[{i}] (source: {c['title']}, page {c['page']})\n{c['text']}"
        for i, c in enumerate(chunks)
    )
    raw = chat(
        [
            {"role": "system", "content": GRADER_SYSTEM},
            {"role": "user", "content": f"Query: {query}\n\nChunks:\n{numbered}"},
        ],
        temperature=0.0,
        max_tokens=256,
    )
    indices, warning = _parse_relevant_indices(raw, len(chunks))
    if indices is None:
        return list(range(len(chunks))), warning
    return indices, None


def rewrite_query(original_query: str, queries_tried: list[str]) -> str:
    """Reformulate the query to improve retrieval when prior attempts failed."""
    history_text = "\n".join(
        f"Attempt {i + 1}: {q}" for i, q in enumerate(queries_tried)
    )
    return chat(
        [
            {
                "role": "system",
                "content": (
                    "You are a query rewriter for a retrieval system. Given an original "
                    "question and the queries already tried (which retrieved no useful "
                    "results), produce a single improved search query that is more likely "
                    "to find relevant documents. Vary the vocabulary and specificity. "
                    "Reply with only the rewritten query — no explanation, no quotes."
                ),
            },
            {
                "role": "user",
                "content": (
                    f"Original question: {original_query}\n\n"
                    f"Queries tried so far:\n{history_text}\n\n"
                    "Rewritten query:"
                ),
            },
        ],
        temperature=0.4,
        max_tokens=128,
    )


ANSWER_SYSTEM = (
    "You are a careful assistant. Answer the user's question using ONLY the provided "
    "context. Cite the document title and page for each claim, like "
    "[Attention Is All You Need, p.3]. Write in clear, plain language that a "
    "non-expert can follow. "
    "If the context does not fully cover the question, say so explicitly rather than "
    "filling the gap from your own knowledge. "
    "Write any mathematical notation as LaTeX wrapped in $...$ for inline math or "
    "$$...$$ for display math. Never use \\( \\) or \\[ \\] delimiters."
)


def answer_stream(query: str, relevant_chunks: list[dict]) -> Iterator[str]:
    context = "\n\n---\n\n".join(
        f"(source: {c['title']}, page {c['page']})\n{c['text']}" for c in relevant_chunks
    )
    yield from chat_stream(
        [
            {"role": "system", "content": ANSWER_SYSTEM},
            {"role": "user", "content": f"Context:\n{context}\n\nQuestion: {query}"},
        ],
        temperature=0.2,
    )


# ---------------------------------------------------------------------------
# Orchestration
# ---------------------------------------------------------------------------

REFUSAL = (
    "I could not find enough relevant information in the indexed documents to answer "
    "that. Rather than guess, I'm declining to answer. Try rephrasing, or ask about "
    "something covered by the corpus."
)


@dataclass
class TraceEvent:
    kind: Literal["retrieve", "grade", "rewrite", "answer", "refuse", "warning"]
    iteration: int
    query: str = ""
    chunks: list[dict] = field(default_factory=list)
    relevant_indices: list[int] = field(default_factory=list)
    message: str = ""


def run_pipeline(
    collection,
    query: str,
    top_k: int = TOP_K,
    max_iterations: int = MAX_RETRIEVAL_ITERATIONS,
    min_relevant: int = MIN_RELEVANT_CHUNKS,
) -> Iterator[TraceEvent]:
    """
    Yields TraceEvents as the loop progresses. The final event is either
    kind="answer" (with chunks set to the grounded context) or kind="refuse".
    The caller streams the answer itself via answer_stream().
    """
    current_query = query
    queries_tried: list[str] = []

    for iteration in range(1, max_iterations + 1):
        chunks = retrieve(collection, current_query, top_k=top_k)
        queries_tried.append(current_query)
        yield TraceEvent("retrieve", iteration, query=current_query, chunks=chunks)

        if not chunks:
            yield TraceEvent("warning", iteration, message="Retrieval returned no chunks.")
            relevant_indices: list[int] = []
        else:
            relevant_indices, warning = grade_chunks(query, chunks)
            if warning:
                yield TraceEvent("warning", iteration, message=warning)
            yield TraceEvent(
                "grade", iteration, query=current_query,
                chunks=chunks, relevant_indices=relevant_indices,
            )

        if len(relevant_indices) >= min_relevant:
            yield TraceEvent(
                "answer", iteration, query=query,
                chunks=[chunks[i] for i in relevant_indices],
            )
            return

        if iteration < max_iterations:
            current_query = rewrite_query(query, queries_tried)
            yield TraceEvent("rewrite", iteration, query=current_query)

    yield TraceEvent("refuse", max_iterations, message=REFUSAL)
