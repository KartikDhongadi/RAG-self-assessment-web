"""
RAG pipeline with a local Ollama model as relevance grader and query rewriter.

Flow: retrieve → grade each chunk → if sufficient context, answer;
      else rewrite query → re-retrieve (up to MAX_RETRIEVAL_ITERATIONS).

Retrieval : ChromaDB with sentence-transformer embeddings (all-MiniLM-L6-v2).
Documents : PDF, Word (.docx), plain text (.txt), Excel (.xlsx/.xls), and CSV
            files found in DOCS_DIR are chunked and indexed automatically.
LLM       : Ollama (local) — change MODEL to any model you have pulled.
"""

import sys
import re
import json
import hashlib
from pathlib import Path

import ollama
import chromadb
from chromadb.utils.embedding_functions import SentenceTransformerEmbeddingFunction
from pypdf import PdfReader

# ── LLM ─────────────────────────────────────────────────────────────────────
# Change to any model you have pulled locally:
#   llama3.2 | mistral | phi3 | gemma2 | qwen2.5 | deepseek-r1
MODEL = "llama3.2"

# ── Pipeline ─────────────────────────────────────────────────────────────────
MAX_RETRIEVAL_ITERATIONS = 3
MIN_RELEVANT_CHUNKS = 1

# ── ChromaDB ─────────────────────────────────────────────────────────────────
CHROMA_PATH = "./chroma_db"
COLLECTION_NAME = "rag_docs_v2"
EMBED_MODEL = "all-MiniLM-L6-v2"

# ── Document ingestion ────────────────────────────────────────────────────────
DOCS_DIR = "./docs"    # folder to scan for documents
CHUNK_SIZE = 1000      # characters per chunk
CHUNK_OVERLAP = 200    # character overlap between consecutive chunks
DOC_INDEX_PATH = "./doc_index.json"  # tracks per-file SHA-256 hashes

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".csv", ".xlsx", ".xls"}


# ---------------------------------------------------------------------------
# PDF loading and chunking
# ---------------------------------------------------------------------------

def _clean(text: str) -> str:
    """Normalize whitespace while preserving sentence structure."""
    text = re.sub(r"[ \t]+", " ", text)          # collapse spaces/tabs
    text = re.sub(r"\n{3,}", "\n\n", text)        # cap consecutive newlines
    text = re.sub(r"(?<!\n)\n(?!\n)", " ", text)  # single newline → space
    return text.strip()


def chunk_text(text: str, chunk_size: int = CHUNK_SIZE, overlap: int = CHUNK_OVERLAP) -> list[str]:
    """Split text into overlapping chunks that always break at sentence boundaries."""
    # Split on sentence-ending punctuation followed by whitespace
    sentences = re.split(r"(?<=[.!?])\s+", text)

    chunks: list[str] = []
    current: list[str] = []
    current_len = 0

    for sentence in sentences:
        s_len = len(sentence)
        if current_len + s_len > chunk_size and current:
            chunks.append(" ".join(current))
            # Roll back by overlap: keep trailing sentences that fit in overlap window
            tail, tail_len = [], 0
            for s in reversed(current):
                if tail_len + len(s) <= overlap:
                    tail.insert(0, s)
                    tail_len += len(s)
                else:
                    break
            current, current_len = tail, tail_len
        current.append(sentence)
        current_len += s_len

    if current:
        chunks.append(" ".join(current))

    return [c for c in chunks if len(c) > 80]   # drop tiny tail chunks


def _make_records(path: Path, sections: list[tuple[int, str]]) -> list[dict]:
    """
    Convert (section_num, text) pairs into chunk record dicts.
    section_num maps to page for PDFs, sheet index for Excel, 1 for flat files.
    """
    records = []
    for section_num, text in sections:
        cleaned = _clean(text)
        if not cleaned:
            continue
        for chunk_idx, chunk in enumerate(chunk_text(cleaned)):
            uid = hashlib.md5(f"{path.name}|{section_num}|{chunk_idx}".encode()).hexdigest()
            records.append({
                "id": uid,
                "text": chunk,
                "metadata": {"source": path.name, "page": section_num, "chunk": chunk_idx},
            })
    return records


def load_pdf(path: Path) -> list[dict]:
    reader = PdfReader(str(path))
    sections = [
        (page_num, page.extract_text() or "")
        for page_num, page in enumerate(reader.pages, start=1)
    ]
    return _make_records(path, sections)


def load_docx(path: Path) -> list[dict]:
    from docx import Document  # python-docx
    doc = Document(str(path))
    lines: list[str] = []
    for para in doc.paragraphs:
        if para.text.strip():
            lines.append(para.text)
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(c.text.strip() for c in row.cells if c.text.strip())
            if row_text:
                lines.append(row_text)
    return _make_records(path, [(1, "\n".join(lines))])


def load_txt(path: Path) -> list[dict]:
    text = path.read_text(encoding="utf-8", errors="replace")
    return _make_records(path, [(1, text)])


def load_excel(path: Path) -> list[dict]:
    import pandas as pd
    xl = pd.ExcelFile(str(path))
    sections: list[tuple[int, str]] = []
    for sheet_idx, sheet_name in enumerate(xl.sheet_names, start=1):
        df = xl.parse(sheet_name).fillna("")
        header = " | ".join(str(c) for c in df.columns)
        rows = [f"Sheet: {sheet_name} | Columns: {header}"]
        for _, row in df.iterrows():
            row_text = " | ".join(
                f"{col}: {val}" for col, val in row.items() if str(val).strip()
            )
            if row_text:
                rows.append(row_text)
        sections.append((sheet_idx, "\n".join(rows)))
    return _make_records(path, sections)


def load_csv(path: Path) -> list[dict]:
    import pandas as pd
    df = pd.read_csv(str(path)).fillna("")
    header = " | ".join(str(c) for c in df.columns)
    rows = [f"Columns: {header}"]
    for _, row in df.iterrows():
        row_text = " | ".join(
            f"{col}: {val}" for col, val in row.items() if str(val).strip()
        )
        if row_text:
            rows.append(row_text)
    return _make_records(path, [(1, "\n".join(rows))])


def load_file(path: Path) -> list[dict]:
    ext = path.suffix.lower()
    if ext == ".pdf":
        return load_pdf(path)
    if ext == ".docx":
        return load_docx(path)
    if ext in {".txt"}:
        return load_txt(path)
    if ext in {".xlsx", ".xls"}:
        return load_excel(path)
    if ext == ".csv":
        return load_csv(path)
    print(f"[Warning] Unsupported file type '{ext}' — skipping {path.name}")
    return []


# ---------------------------------------------------------------------------
# ChromaDB setup
# ---------------------------------------------------------------------------

def _file_hash(path: Path) -> str:
    """Return the SHA-256 hex digest of a file's contents."""
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for block in iter(lambda: f.read(65536), b""):
            h.update(block)
    return h.hexdigest()


def _load_doc_index() -> dict[str, str]:
    """Load the persisted {filename: sha256} index, or return empty dict."""
    index_path = Path(DOC_INDEX_PATH)
    if index_path.exists():
        return json.loads(index_path.read_text())
    return {}


def _save_doc_index(index: dict[str, str]) -> None:
    Path(DOC_INDEX_PATH).write_text(json.dumps(index, indent=2))


def _index_records(collection: chromadb.Collection, records: list[dict]) -> None:
    """Add records to the collection in batches of 500."""
    batch_size = 500
    for i in range(0, len(records), batch_size):
        batch = records[i : i + batch_size]
        collection.add(
            ids=[r["id"] for r in batch],
            documents=[r["text"] for r in batch],
            metadatas=[r["metadata"] for r in batch],
        )


def setup_chroma() -> chromadb.Collection:
    """
    Initialise a persistent ChromaDB collection.
    Scans DOCS_DIR for supported file types and:
      - indexes new files that haven't been seen before
      - re-indexes files whose content has changed since last run
      - skips files that are unchanged
    Safe to call on every run.
    """
    embedding_fn = SentenceTransformerEmbeddingFunction(model_name=EMBED_MODEL)
    db = chromadb.PersistentClient(path=CHROMA_PATH)
    collection = db.get_or_create_collection(
        name=COLLECTION_NAME,
        embedding_function=embedding_fn,
        metadata={"hnsw:space": "cosine"},
    )

    docs_path = Path(DOCS_DIR)
    doc_files = sorted(
        f for f in docs_path.iterdir()
        if f.is_file() and f.suffix.lower() in SUPPORTED_EXTENSIONS
    )
    if not doc_files:
        exts = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        print(f"[ChromaDB] No supported files found in '{DOCS_DIR}'. "
              f"Supported: {exts}")
        sys.exit(1)

    stored_index = _load_doc_index()
    updated_index = dict(stored_index)

    new_count = updated_count = skipped_count = 0

    for doc_path in doc_files:
        current_hash = _file_hash(doc_path)
        stored_hash = stored_index.get(doc_path.name)

        if stored_hash == current_hash:
            skipped_count += 1
            print(f"[Docs] {doc_path.name} — unchanged, skipping")
            continue

        records = load_file(doc_path)
        if not records:
            continue

        if stored_hash is None:
            print(f"[Docs] {doc_path.name} — new ({len(records)} chunks), indexing...")
            _index_records(collection, records)
            new_count += 1
        else:
            print(f"[Docs] {doc_path.name} — updated ({len(records)} chunks), re-indexing...")
            collection.delete(where={"source": doc_path.name})
            _index_records(collection, records)
            updated_count += 1

        updated_index[doc_path.name] = current_hash

    _save_doc_index(updated_index)

    print(
        f"\n[ChromaDB] Done — {new_count} new, {updated_count} updated, "
        f"{skipped_count} unchanged. "
        f"Collection '{COLLECTION_NAME}' has {collection.count()} chunks.\n"
    )
    return collection


def make_chroma_retriever(collection: chromadb.Collection, top_k: int = 5):
    """Return a retriever function bound to the given ChromaDB collection."""
    def retrieve(query: str) -> list[str]:
        results = collection.query(query_texts=[query], n_results=top_k)
        chunks = results["documents"][0]
        metas = results["metadatas"][0]
        # Print source attribution so the demo shows where answers come from
        for chunk, meta in zip(chunks, metas):
            print(f"    [{meta['source']}  p.{meta['page']}]  {chunk[:80]}...")
        return chunks
    return retrieve


# ---------------------------------------------------------------------------
# Core pipeline functions
# ---------------------------------------------------------------------------

def grade_chunk(query: str, chunk: str) -> bool:
    """Return True if the chunk is relevant to the query."""
    response = ollama.chat(
        model=MODEL,
        messages=[
            {
                "role": "system",
                "content": (
                    "You are a relevance grader. Given a user query and a document chunk, "
                    "decide whether the chunk contains information useful for answering the query. "
                    "Reply with exactly one word: YES or NO. No explanation."
                ),
            },
            {
                "role": "user",
                "content": f"Query: {query}\n\nChunk:\n{chunk}",
            },
        ],
    )
    return response.message.content.strip().upper().startswith("YES")


def rewrite_query(original_query: str, queries_tried: list[str]) -> str:
    """Reformulate the query to improve retrieval when prior attempts failed."""
    history_text = "\n".join(
        f"Attempt {i + 1}: {q}" for i, q in enumerate(queries_tried)
    )
    response = ollama.chat(
        model=MODEL,
        messages=[
            {
                "role": "system",
                "content": (
                    "You are a query rewriter for a retrieval system. "
                    "Given an original question and the queries already tried (which retrieved no useful results), "
                    "produce a single improved search query that is more likely to find relevant documents. "
                    "Reply with only the rewritten query — no explanation, no quotes."
                ),
            },
            {
                "role": "user",
                "content": (
                    f"Original question: {original_query}\n\n"
                    f"Queries tried so far:\n{history_text}\n\n"
                    "Rewritten query:"
                ),
            },
        ],
    )
    return response.message.content.strip()


def generate_answer(query: str, relevant_chunks: list[str]) -> str:
    """Stream an answer grounded in the provided relevant chunks."""
    context = "\n\n---\n\n".join(relevant_chunks)
    print("\n[Generating answer...]\n")

    stream = ollama.chat(
        model=MODEL,
        messages=[
            {
                "role": "system",
                "content": (
                    "You are a helpful assistant. Answer the user's question using only "
                    "the provided context. If the context does not fully cover the question, "
                    "say so honestly."
                ),
            },
            {
                "role": "user",
                "content": f"Context:\n{context}\n\nQuestion: {query}",
            },
        ],
        stream=True,
    )

    full_answer = []
    for chunk in stream:
        token = chunk.message.content
        print(token, end="", flush=True)
        full_answer.append(token)

    print()
    return "".join(full_answer)


# ---------------------------------------------------------------------------
# Main pipeline loop
# ---------------------------------------------------------------------------

def run_rag_pipeline(query: str, retriever) -> str:
    """
    Orchestrates: retrieve → grade → answer (or rewrite → repeat).

    Args:
        query:     The user's original question.
        retriever: Callable(query: str) -> list[str] returning text chunks.

    Returns:
        A grounded answer string, or a fallback message if no context was found.
    """
    current_query = query
    queries_tried: list[str] = []

    for iteration in range(1, MAX_RETRIEVAL_ITERATIONS + 1):
        print(f"\n[Iteration {iteration}] Retrieving for: {current_query!r}")
        chunks = retriever(current_query)
        queries_tried.append(current_query)

        print(f"\n  Grading {len(chunks)} chunk(s) for relevance...")
        relevant = [chunk for chunk in chunks if grade_chunk(query, chunk)]
        print(f"  {len(relevant)}/{len(chunks)} chunk(s) deemed relevant.")

        if len(relevant) >= MIN_RELEVANT_CHUNKS:
            return generate_answer(query, relevant)

        if iteration < MAX_RETRIEVAL_ITERATIONS:
            print("  Insufficient context — rewriting query...")
            current_query = rewrite_query(query, queries_tried)
        else:
            print("  Max iterations reached with no sufficient context.")

    return (
        "I was unable to find enough relevant information to answer your question. "
        "Please try rephrasing or providing more context."
    )


# ---------------------------------------------------------------------------
# Entry point
# ---------------------------------------------------------------------------

if __name__ == "__main__":
    print(f"Model : {MODEL}")
    print("Setting up ChromaDB...\n")

    chroma_collection = setup_chroma()
    retriever = make_chroma_retriever(chroma_collection, top_k=5)

    print("RAG pipeline ready. Type your question below (or 'quit' to exit).\n")

    while True:
        try:
            user_query = input("You: ").strip()
        except (EOFError, KeyboardInterrupt):
            print("\nGoodbye!")
            break

        if not user_query:
            continue
        if user_query.lower() in {"quit", "exit", "q"}:
            print("Goodbye!")
            break

        answer = run_rag_pipeline(user_query, retriever)
        print(f"\nAnswer: {answer}\n")
        print("-" * 60 + "\n")
